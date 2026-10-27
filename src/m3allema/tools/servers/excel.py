"""Excel MCP server: read and write ``.xlsx`` workbooks inside the files folder.

Run standalone::

    M3ALLEMA_FILES_ROOT=./files python -m m3allema.tools.servers.excel
"""

from __future__ import annotations

import csv
import io
import os
from collections.abc import Sequence
from pathlib import Path
from typing import Any

from mcp.server.mcpserver import MCPServer
from openpyxl import Workbook, load_workbook

from m3allema.tools.servers._annotations import LOCAL_WRITE, READ_ONLY
from m3allema.tools.servers.files import resolve_inside

MAX_ROWS = 500

Cell = str | int | float | bool | None


def _require_xlsx(path: Path) -> None:
    if path.suffix.lower() != ".xlsx":
        raise ValueError("only .xlsx workbooks are supported")


def rows_to_csv(rows: Sequence[Sequence[Any]]) -> str:
    """Render rows as CSV text; ``None`` cells become empty."""
    buffer = io.StringIO()
    writer = csv.writer(buffer, lineterminator="\n")
    for row in rows:
        writer.writerow(["" if value is None else value for value in row])
    return buffer.getvalue().rstrip("\n")


def create_excel_server(root: Path) -> MCPServer:
    """Build an Excel server confined to ``root``."""
    root = root.resolve()
    server = MCPServer(
        "excel", instructions="Read and write .xlsx workbooks in the business's files folder."
    )

    @server.tool(annotations=READ_ONLY)
    def list_sheets(path: str) -> str:
        """List the sheets of a workbook with their sizes."""
        target = resolve_inside(root, path)
        _require_xlsx(target)
        workbook = load_workbook(target, read_only=True, data_only=True)
        try:
            return "\n".join(
                f"{ws.title}: {ws.max_row} rows x {ws.max_column} columns"
                for ws in workbook.worksheets
            )
        finally:
            workbook.close()

    @server.tool(annotations=READ_ONLY)
    def read_sheet(path: str, sheet: str | None = None, start_row: int = 1) -> str:
        """Read a sheet as CSV (computed values, not formulas), up to 500 rows per call."""
        target = resolve_inside(root, path)
        _require_xlsx(target)
        workbook = load_workbook(target, read_only=True, data_only=True)
        try:
            ws = workbook[sheet] if sheet else workbook.worksheets[0]
            first = max(start_row, 1)
            rows = list(ws.iter_rows(min_row=first, max_row=first + MAX_ROWS - 1, values_only=True))
            text = rows_to_csv(rows) or "(empty sheet)"
            last = first + len(rows) - 1
            if ws.max_row and last < ws.max_row:
                text += (
                    f"\n\n[rows {first}-{last} of {ws.max_row}; continue with start_row={last + 1}]"
                )
            return text
        except KeyError as exc:
            raise ValueError(f"no sheet named {sheet!r}") from exc
        finally:
            workbook.close()

    @server.tool(annotations=LOCAL_WRITE)
    def write_sheet(
        path: str, rows: list[list[Cell]], sheet: str = "Sheet1", overwrite: bool = False
    ) -> str:
        """Create a workbook with one sheet from rows (the first row is usually the header).

        Set overwrite=true to replace an existing file.
        """
        target = resolve_inside(root, path)
        _require_xlsx(target)
        if target.exists() and not overwrite:
            raise ValueError(f"{path!r} already exists; pass overwrite=true to replace it")
        workbook = Workbook()
        ws = workbook.active
        assert ws is not None  # noqa: S101 - a new workbook always has an active sheet
        ws.title = sheet
        for row in rows:
            ws.append(row)
        target.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(target)
        return f"Wrote {len(rows)} rows to {path} ({sheet})"

    @server.tool(annotations=LOCAL_WRITE)
    def append_rows(path: str, rows: list[list[Cell]], sheet: str | None = None) -> str:
        """Add rows to the end of a sheet in an existing workbook."""
        target = resolve_inside(root, path)
        _require_xlsx(target)
        if not target.is_file():
            raise ValueError(f"{path!r} does not exist")
        workbook = load_workbook(target)
        try:
            ws = workbook[sheet] if sheet else workbook.worksheets[0]
        except KeyError as exc:
            raise ValueError(f"no sheet named {sheet!r}") from exc
        for row in rows:
            ws.append(row)
        workbook.save(target)
        return f"Appended {len(rows)} rows to {path} ({ws.title})"

    return server


if __name__ == "__main__":
    create_excel_server(Path(os.environ.get("M3ALLEMA_FILES_ROOT", "."))).run("stdio")
