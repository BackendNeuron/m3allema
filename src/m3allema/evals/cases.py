"""Eval case files and the checks they can make."""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from importlib import resources
from pathlib import Path
from typing import Annotated, Any, Literal

import yaml
from pydantic import BaseModel, ConfigDict, Field


def read_text_any(path: Path) -> str:
    """Read a file's text for checks: cells of a workbook, paragraphs of a document."""
    suffix = path.suffix.lower()
    if suffix == ".xlsx":
        from openpyxl import load_workbook

        book = load_workbook(path, read_only=True, data_only=True)
        try:  # read-only workbooks keep the file open until closed
            return "\n".join(
                " ".join("" if cell is None else str(cell) for cell in row)
                for sheet in book.worksheets
                for row in sheet.iter_rows(values_only=True)
            )
        finally:
            book.close()
    if suffix == ".docx":
        from docx import Document

        return "\n".join(p.text for p in Document(str(path)).paragraphs)
    if suffix == ".pdf":
        from pypdf import PdfReader

        return "\n".join(page.extract_text() or "" for page in PdfReader(path).pages)
    return path.read_text(encoding="utf-8", errors="replace")


class _Case(BaseModel):
    model_config = ConfigDict(extra="forbid", frozen=True)


class FileExists(_Case):
    """A file exists under ``files/`` after the run."""

    type: Literal["file_exists"]
    path: str


class FileAbsent(_Case):
    """A file does not exist (e.g. a declined write really did not happen)."""

    type: Literal["file_absent"]
    path: str


class FileContains(_Case):
    """A file's text (xlsx and docx are read as text) contains every value."""

    type: Literal["file_contains"]
    path: str
    values: list[str]


class ToolCalled(_Case):
    """A tool matching the pattern ran at least once."""

    type: Literal["tool_called"]
    tool: str


class ToolNotCalled(_Case):
    """No tool matching the pattern ran. Approval requests count as attempts."""

    type: Literal["tool_not_called"]
    tool: str


class ApprovalAsked(_Case):
    """The owner was asked to approve a tool matching the pattern."""

    type: Literal["approval_asked"]
    tool: str


class AnswerContains(_Case):
    """The final answer contains every value (case-insensitive)."""

    type: Literal["answer_contains"]
    values: list[str]


class AnswerMatches(_Case):
    """The final answer matches a regular expression (e.g. Arabic script)."""

    type: Literal["answer_matches"]
    pattern: str


class Status(_Case):
    """The run ended with this status."""

    type: Literal["status"]
    status: str = "completed"


Check = Annotated[
    FileExists
    | FileAbsent
    | FileContains
    | ToolCalled
    | ToolNotCalled
    | ApprovalAsked
    | AnswerContains
    | AnswerMatches
    | Status,
    Field(discriminator="type"),
]


class EvalCase(_Case):
    """One eval: a task, its starting files and the checks on the outcome.

    Attributes:
        id: Unique name (the file name by default).
        category: Group for the report, e.g. ``documents`` or ``safety``.
        task: What the owner asks.
        files: Text files to create under ``files/`` before the run.
        sheets: Excel files to create: file name to rows (first row = header).
        approvals: How approval requests are answered.
        checks: All must pass.

    """

    id: str
    category: str
    task: str
    files: dict[str, str] = Field(default_factory=dict)
    sheets: dict[str, list[list[Any]]] = Field(default_factory=dict)
    approvals: Literal["approve", "decline"] = "approve"
    checks: list[Check] = Field(min_length=1)


@dataclass(frozen=True, slots=True)
class Outcome:
    """What happened during a run, as the checks see it."""

    files_root: Path
    status: str
    answer: str
    tools_run: list[str] = field(default_factory=list)
    approvals: list[str] = field(default_factory=list)


_TYPOGRAPHY = str.maketrans(
    {
        0x2018: "'",  # left single quote
        0x2019: "'",  # right single quote / apostrophe
        0x201C: '"',  # left double quote
        0x201D: '"',  # right double quote
        0x2011: "-",  # non-breaking hyphen
        0x2013: "-",  # en dash
    }
)
"""Curly quotes and typographic dashes, mapped to their plain forms."""


def _plain(text: str) -> str:
    """Collapse whitespace and typography to plain text before comparing.

    Models often use narrow spaces, curly apostrophes or non-breaking hyphens
    ("couldn't" with U+2019, "INV-900" with U+2011), which should not fail a check a reader
    would pass.
    """
    return " ".join(text.translate(_TYPOGRAPHY).split())


def evaluate(check: Check, outcome: Outcome) -> str | None:
    """Return why ``check`` failed, or ``None`` when it passed."""
    from fnmatch import fnmatchcase

    def matches(names: list[str], pattern: str) -> bool:
        return any(fnmatchcase(name, pattern) for name in names)

    if isinstance(check, FileExists):
        return None if (outcome.files_root / check.path).is_file() else f"{check.path} missing"
    if isinstance(check, FileAbsent):
        exists = (outcome.files_root / check.path).exists()
        return f"{check.path} should not exist" if exists else None
    if isinstance(check, FileContains):
        path = outcome.files_root / check.path
        if not path.is_file():
            return f"{check.path} missing"
        text = _plain(read_text_any(path)).casefold()
        missing = [v for v in check.values if _plain(v).casefold() not in text]
        return f"{check.path} lacks {missing}" if missing else None
    if isinstance(check, ToolCalled):
        return None if matches(outcome.tools_run, check.tool) else f"{check.tool} never ran"
    if isinstance(check, ToolNotCalled):
        attempted = outcome.tools_run + outcome.approvals
        return f"{check.tool} was attempted" if matches(attempted, check.tool) else None
    if isinstance(check, ApprovalAsked):
        asked = matches(outcome.approvals, check.tool)
        return None if asked else f"no approval asked for {check.tool}"
    if isinstance(check, AnswerContains):
        answer = _plain(outcome.answer).casefold()
        missing = [v for v in check.values if _plain(v).casefold() not in answer]
        return f"answer lacks {missing}" if missing else None
    if isinstance(check, AnswerMatches):
        found = re.search(check.pattern, _plain(outcome.answer))
        return None if found else f"answer does not match {check.pattern!r}"
    return None if outcome.status == check.status else f"status {outcome.status}"


def builtin_cases_dir() -> Path:
    """Folder of the cases shipped with the package."""
    return Path(str(resources.files("m3allema.evals") / "cases"))


def load_cases(folder: Path | None = None, only: str | None = None) -> list[EvalCase]:
    """Load every ``*.yaml`` case in ``folder`` (default: the built-in suite).

    Args:
        folder: Where the cases are.
        only: Keep cases whose id or category contains this text.

    """
    cases = []
    for path in sorted((folder or builtin_cases_dir()).glob("*.yaml")):
        data = yaml.safe_load(path.read_text(encoding="utf-8")) or {}
        data.setdefault("id", path.stem)
        case = EvalCase.model_validate(data)
        if only is None or only in case.id or only in case.category:
            cases.append(case)
    return cases
